import docx

doc = docx.Document()
doc.add_heading("A Test Manuscript", 1)
doc.add_paragraph("This is a paragraph with some normal text.")
doc.add_paragraph("Another paragraph with [1] citation.")

doc.save("test_manuscript.docx")
print("Saved dummy manuscript to test_manuscript.docx")
